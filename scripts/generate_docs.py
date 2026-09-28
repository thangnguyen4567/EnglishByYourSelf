"""Sinh nội dung docs + dữ liệu JSON từ file Excel ngữ pháp.

Chạy:  python -X utf8 scripts/generate_docs.py [đường_dẫn_xlsx]
Mặc định đọc ../Ngu-phap-tieng-Anh-30-ngay.xlsx (thư mục cha của website/).

Các file sinh ra (sẽ bị ghi đè mỗi lần chạy):
  src/data/*.json              dữ liệu cho component React (Quiz, tiến độ...)
  docs/ngu-phap/**             29 trang chủ đề, nhóm theo phân loại ngữ pháp
  docs/lo-trinh/tuan-*.mdx     kế hoạch từng tuần
  docs/bai-kiem-tra/*.mdx      5 bài kiểm tra tương tác
  docs/tra-cuu/*.mdx           bảng tra cứu nhanh, sổ lỗi thường gặp
Các trang viết tay (intro, tong-hop-12-thi, lo-trinh/index...) không bị động tới.
"""
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT.parent / "Ngu-phap-tieng-Anh-30-ngay.xlsx"
DOCS = ROOT / "docs"
DATA = ROOT / "src" / "data"

# ---------------------------------------------------------------------------
# Phân loại khoa học: nhóm lớn -> (nhóm con) -> chủ đề.
# Thứ tự đi từ đơn vị nhỏ (từ loại) đến câu đơn, câu phức và cấu trúc nâng cao.
# ---------------------------------------------------------------------------
GROUPS = [
    {"key": "Nền tảng", "dir": "01-nen-tang", "label": "1. Nền tảng",
     "desc": "Từ loại và cấu trúc câu cơ bản – nền móng cho mọi chủ đề khác."},
    {"key": "Thì", "dir": "02-thi", "label": "2. Các thì (Tenses)",
     "desc": "Hệ thống thì theo trục thời gian: hiện tại → quá khứ → tương lai.",
     "sub": [
         {"dir": "01-hien-tai", "label": "Nhóm hiện tại", "codes": ["G02", "G03", "G06", "G07"]},
         {"dir": "02-qua-khu", "label": "Nhóm quá khứ", "codes": ["G04", "G05", "G08"]},
         {"dir": "03-tuong-lai", "label": "Nhóm tương lai", "codes": ["G09", "G10"]},
     ]},
    {"key": "Động từ", "dir": "03-dong-tu", "label": "3. Động từ",
     "desc": "Động từ khuyết thiếu và dạng của động từ theo sau (V-ing / to V)."},
    {"key": "Danh từ", "dir": "04-danh-tu", "label": "4. Danh từ & Đại từ",
     "desc": "Danh từ đếm được/không đếm được, mạo từ, lượng từ, đại từ."},
    {"key": "Tính/Trạng từ", "dir": "05-tinh-tu-trang-tu", "label": "5. Tính từ & Trạng từ",
     "desc": "Vị trí, trật tự tính từ và các dạng so sánh."},
    {"key": "Giới từ", "dir": "06-gioi-tu", "label": "6. Giới từ",
     "desc": "Giới từ thời gian, nơi chốn và cụm tính từ + giới từ cố định."},
    {"key": "Cấu trúc câu", "dir": "07-cau-truc-cau", "label": "7. Cấu trúc câu",
     "desc": "Câu bị động và các dạng câu hỏi."},
    {"key": "Câu phức", "dir": "08-cau-phuc", "label": "8. Câu phức",
     "desc": "Câu điều kiện, câu ước, mệnh đề quan hệ, câu tường thuật, liên từ."},
    {"key": "Nâng cao", "dir": "09-nang-cao", "label": "9. Cấu trúc nâng cao",
     "desc": "Used to, thể sai khiến, đảo ngữ, hòa hợp chủ ngữ – động từ."},
]

TEST_SHEETS = [
    ("tuan-1", "Test Tuần 1"),
    ("tuan-2", "Test Tuần 2"),
    ("tuan-3", "Test Tuần 3"),
    ("tuan-4", "Test Tuần 4"),
    ("cuoi-khoa", "Test Cuối Khóa"),
]

# Chủ đề trong Excel được tách thành nhiều trang (cùng mã, dùng chung cho bài kiểm tra & lộ trình).
# "pick" chọn dòng dữ liệu Excel (theo chỉ số dòng) cho từng trang; "all" = lấy hết, [] = không lấy.
# Nội dung bổ sung: content/grammar/<key>.mdx · Bài kiểm tra cấp độ: content/quizzes/<key>.json
SPLIT_TOPICS = {
    "G01": [
        {"key": "g01a", "label": "G01a", "slug": "g01-tu-loai",
         "title": "Từ loại (Parts of speech)", "short": "Từ loại",
         "pick": {"structure": [1], "usage": "all", "signals": "all", "examples": [], "mistakes": "all"}},
        {"key": "g01b", "label": "G01b", "slug": "g01-cau-truc-cau-co-ban",
         "title": "Cấu trúc câu cơ bản (Basic sentence structure)", "short": "Cấu trúc câu cơ bản",
         "pick": {"structure": [0], "usage": [], "signals": [], "examples": [0, 1], "mistakes": []}},
    ],
}

SIGN_LABELS = {"(+)": "Khẳng định", "(–)": "Phủ định", "(-)": "Phủ định", "(?)": "Nghi vấn"}


def slugify(text: str) -> str:
    text = text.replace("đ", "d").replace("Đ", "D")
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    text = re.sub(r"\(.*?\)", "", text)
    text = re.sub(r"[^a-zA-Z0-9]+", "-", text).strip("-").lower()
    return re.sub(r"-+", "-", text)


def esc(text: str) -> str:
    """Escape ký tự đặc biệt của MDX trong văn bản thường."""
    return (text.replace("\\", "\\\\").replace("{", "\\{").replace("}", "\\}")
            .replace("<", "&lt;").replace(">", "&gt;").replace("*", "\\*").replace("_", "\\_"))


def esc_cell(text: str) -> str:
    return esc(text).replace("|", "\\|").replace("\n", "<br/>")


def lines(value) -> list[str]:
    return [l.strip() for l in str(value or "").split("\n") if l.strip()]


def short_title(title: str) -> str:
    """'Hiện tại đơn (Present Simple)' -> 'Hiện tại đơn'."""
    return re.sub(r"\s*\(.*?\)\s*$", "", title).strip()


def english_name(title: str) -> str | None:
    m = re.search(r"\(([^)]*[A-Za-z][^)]*)\)\s*$", title)
    return m.group(1) if m else None


# ---------------------------------------------------------------------------
# Đọc Excel
# ---------------------------------------------------------------------------
wb = openpyxl.load_workbook(XLSX)

topics = []
for r in wb["Ngữ pháp"].iter_rows(min_row=5, values_only=True):
    if not r[0]:
        continue
    code, group, title, structure, usage, signals, examples, mistakes = r[:8]
    topics.append({
        "code": code, "group": group, "title": title, "short": short_title(title),
        "en": english_name(title), "structure": lines(structure), "usage": lines(usage),
        "signals": lines(signals), "examples": lines(examples), "mistakes": lines(mistakes),
    })
by_code = {t["code"]: t for t in topics}

roadmap = []
for r in wb["Lộ trình 30 ngày"].iter_rows(min_row=6, values_only=True):
    if not r[0]:
        continue
    day, week, _, topic, ref, tasks, minutes = r[:7]
    codes = re.findall(r"G\d{2}", ref or "")
    if "–" in (ref or ""):  # dạng khoảng G01–G07
        a, b = [int(x[1:]) for x in re.findall(r"G\d{2}", ref)]
        codes = [f"G{i:02d}" for i in range(a, b + 1)]
    roadmap.append({
        "day": day, "week": week, "topic": topic, "ref": ref, "codes": codes,
        "minutes": minutes, "isTest": "TEST" in topic.upper(),
        "tasks": [t.strip() for t in tasks.split("|")],
    })

answers = {}
ans_ws = wb["Đáp án"]
for r in ans_ws.iter_rows(min_row=2, values_only=True):
    if r[0]:
        answers[(r[0], r[1])] = (r[2], r[3])

# Câu hỏi bổ sung (Trung bình + Khó) viết tay cho bài kiểm tra tuần: content/tests/<id>.json
#   {"normal": [câu hỏi], "hard": [câu hỏi]}; câu hỏi = {"code","q","options","answer","explain"}
# Được nối sau các câu "Cơ bản" lấy từ Excel.
TEST_EXTRA = ROOT / "content" / "tests"


def load_test_extra(test_id: str, subtitle: str, start: int) -> list[dict]:
    path = TEST_EXTRA / f"{test_id}.json"
    if not path.exists():
        return []
    lo, hi = (int(x[1:]) for x in re.findall(r"G\d{2}", subtitle)[:2])
    data = json.loads(path.read_text(encoding="utf-8"))
    out, errors = [], []
    for level in ("normal", "hard"):
        for i, q in enumerate(data.get(level, []), start=1):
            where = f"{level} #{i}"
            code = q.get("code", "")
            if not re.fullmatch(r"G\d{2}", code) or not lo <= int(code[1:]) <= hi:
                errors.append(f"{where}: code '{code}' ngoài phạm vi G{lo:02d}–G{hi:02d}")
            opts = q.get("options", {})
            if sorted(opts) != list("ABCD") or len({str(v).strip().lower() for v in opts.values()}) != 4:
                errors.append(f"{where}: options phải có đủ A-D, không trùng")
            if q.get("answer") not in ("A", "B", "C", "D"):
                errors.append(f"{where}: answer phải là A/B/C/D")
            if not q.get("q") or not q.get("explain"):
                errors.append(f"{where}: thiếu câu hỏi hoặc giải thích")
            out.append({"no": start + len(out) + 1, "code": code, "q": q.get("q", ""), "level": level,
                        "options": {k: str(v) for k, v in opts.items()},
                        "answer": q.get("answer"), "explain": q.get("explain", "")})
    if errors:
        raise SystemExit(f"Lỗi trong content/tests/{path.name}:\n  " + "\n  ".join(errors))
    return out


tests = []
for test_id, sheet in TEST_SHEETS:
    ws = wb[sheet]
    heading = ws["A1"].value  # "TEST TUẦN 1 – Các thì cơ bản (G01–G07)"
    subtitle = heading.split("–", 1)[1].strip()
    day = int(re.search(r"Ngày (\d+)", ws["A2"].value).group(1))
    questions = []
    for r in ws.iter_rows(min_row=6, values_only=True):
        if r[0] is None:
            continue
        ans, explain = answers[(sheet, r[0])]
        questions.append({
            "no": r[0], "code": r[1], "q": r[2], "level": "basic",
            "options": {k: str(v) for k, v in zip("ABCD", r[3:7])},
            "answer": ans, "explain": explain,
        })
    questions += load_test_extra(test_id, subtitle, len(questions))
    name = sheet.replace("Test ", "Bài kiểm tra ")
    tests.append({"id": test_id, "name": name, "subtitle": subtitle, "day": day,
                  "questions": questions})

# ---------------------------------------------------------------------------
# Đường dẫn cho từng chủ đề
# ---------------------------------------------------------------------------
topic_day = {}
for d in roadmap:
    if not d["isTest"]:
        for c in d["codes"]:
            topic_day.setdefault(c, d["day"])

for g in GROUPS:
    members = [t for t in topics if t["group"] == g["key"]]
    for t in members:
        sub = next((s for s in g.get("sub", []) if t["code"] in s["codes"]), None)
        rel = f"ngu-phap/{g['dir']}" + (f"/{sub['dir']}" if sub else "")
        t["slug"] = f"{t['code'].lower()}-{slugify(t['short'])}"
        t["dir"] = rel
        t["path"] = f"/docs/ngu-phap/{t['slug']}"
        t["day"] = topic_day.get(t["code"])
        t["groupLabel"] = g["label"]

def expand(t: dict) -> list[dict]:
    """Một chủ đề Excel → 1 trang, hoặc nhiều trang nếu có trong SPLIT_TOPICS."""
    if t["code"] not in SPLIT_TOPICS:
        return [{**t, "key": t["code"].lower(), "label": t["code"]}]
    out = []
    for part in SPLIT_TOPICS[t["code"]]:
        page = {**t, **{k: part[k] for k in ("key", "label", "slug", "title", "short")}}
        page["path"] = f"/docs/ngu-phap/{part['slug']}"
        for field, sel in part["pick"].items():
            page[field] = t[field] if sel == "all" else [t[field][i] for i in sel if i < len(t[field])]
        out.append(page)
    return out


missing = [t["code"] for t in topics if "slug" not in t]
assert not missing, f"Chủ đề chưa được xếp nhóm: {missing}"
pages = [p for t in topics for p in expand(t)]
for t in topics:  # đường dẫn chính của một mã = trang đầu tiên của mã đó
    t["path"] = expand(t)[0]["path"]
pages_by_code = {}
for pg in pages:
    pages_by_code.setdefault(pg["code"], []).append(pg)

# ---------------------------------------------------------------------------
# Ghi JSON
# ---------------------------------------------------------------------------
DATA.mkdir(parents=True, exist_ok=True)


def dump(name, obj):
    (DATA / name).write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")


dump("topics.json", {t["code"]: {k: t[k] for k in ("code", "short", "title", "group", "path", "day")}
                     for t in topics})
dump("tests.json", tests)
dump("pages.json", [{k: pg[k] for k in ("code", "label", "short", "path")} for pg in pages])
# Số câu ghi trong nhiệm vụ ngày kiểm tra (Excel) -> số câu thực tế sau khi bổ sung
for d in roadmap:
    test = next((t for t in tests if t["day"] == d["day"]), None)
    if d["isTest"] and test:
        d["tasks"] = [re.sub(r"\(\d+ câu", f"({len(test['questions'])} câu", x) for x in d["tasks"]]
dump("roadmap.json", roadmap)

# ---------------------------------------------------------------------------
# Sinh trang chủ đề ngữ pháp
# ---------------------------------------------------------------------------
GEN_DIR = DOCS / "ngu-phap"
for g in GROUPS:  # chỉ xoá file sinh tự động, giữ trang viết tay
    for f in (GEN_DIR / g["dir"]).glob("**/g[0-9][0-9]-*.mdx"):
        f.unlink()


def write(path: Path, content: str):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content.strip() + "\n", encoding="utf-8")


def category(path: Path, label: str, position: int, desc: str | None = None, collapsed=True):
    cat = {"label": label, "position": position, "collapsed": collapsed}
    if desc:
        cat["link"] = {"type": "generated-index", "description": desc}
    write(path / "_category_.json", json.dumps(cat, ensure_ascii=False, indent=2))


def labeled_list(items: list[str]) -> str:
    """Dòng dạng 'nhãn: nội dung' -> in đậm nhãn."""
    out = []
    for line in items:
        m = re.match(r"^([^:]{1,45}):\s*(.+)$", line)
        if m and not m.group(1).startswith(("http", "✗")):
            out.append(f"- **{esc(m.group(1))}:** {esc(m.group(2))}")
        else:
            out.append(f"- {esc(line)}")
    return "\n".join(out)


def structure_block(items: list[str]) -> str:
    signed = [l for l in items if l[:3] in SIGN_LABELS]
    if signed and len(signed) == len(items):
        rows = "\n".join(f"| {SIGN_LABELS[l[:3]]} `{l[:3]}` | `{l[3:].strip()}` |" for l in items)
        return f"| Dạng câu | Công thức |\n|---|---|\n{rows}"
    out = []
    for line in items:
        m = re.match(r"^([^:]{1,45}):\s*(.+)$", line)
        if m:
            out.append(f"- **{esc(m.group(1))}:** `{m.group(2)}`")
        else:
            out.append(f"- `{line}`")
    return "\n".join(out)


def mistakes_block(items: list[str]) -> str:
    pairs, notes = [], []
    for line in items:
        if "✗" in line and "→" in line:
            wrong, right = line.split("→", 1)
            pairs.append((wrong.replace("✗", "").strip(), right.replace("✓", "").strip()))
        else:
            notes.append(line)
    out = []
    if pairs:
        out.append("| ❌ Sai | ✅ Đúng |\n|---|---|")
        out += [f"| {esc_cell(w)} | {esc_cell(r)} |" for w, r in pairs]
    if notes:
        out.append("\n:::tip[Ghi nhớ]\n\n" + "\n".join(f"- {esc(n)}" for n in notes) + "\n\n:::")
    return "\n".join(out)


# ---------------------------------------------------------------------------
# Nội dung bổ sung viết tay: content/grammar/gXX.mdx
# File chia khối bằng dòng đánh dấu {/* @ten-khoi */}; mỗi khối được chèn vào trang:
#   intro     -> đầu trang (trước mục Cấu trúc)
#   structure -> cuối mục Cấu trúc        usage    -> cuối mục Cách dùng
#   signals   -> cuối mục Dấu hiệu        examples -> cuối mục Ví dụ
#   mistakes  -> cuối mục Lỗi thường gặp
#   compare   -> mục riêng "Phân biệt & mở rộng"
#   summary   -> mục riêng "Tóm tắt ghi nhớ"
# ---------------------------------------------------------------------------
ENRICH = ROOT / "content" / "grammar"
ENRICH_KEYS = {"intro", "structure", "usage", "signals", "examples", "mistakes", "compare", "summary"}
MARKER = re.compile(r"^\{/\*\s*@(\w+)\s*\*/\}\s*$")


def load_enrichment(key: str) -> dict[str, str]:
    path = ENRICH / f"{key}.mdx"
    if not path.exists():
        return {}
    blocks, key, buf = {}, None, []
    for line in path.read_text(encoding="utf-8").splitlines():
        m = MARKER.match(line)
        if m:
            if key:
                blocks[key] = "\n".join(buf).strip()
            key, buf = m.group(1), []
            assert key in ENRICH_KEYS, f"{path.name}: khối không hợp lệ @{key}"
        elif key:
            buf.append(line)
    if key:
        blocks[key] = "\n".join(buf).strip()
    return {k: v for k, v in blocks.items() if v}


# ---------------------------------------------------------------------------
# Bài kiểm tra 3 cấp độ viết tay: content/quizzes/gXX.json
#   {"easy": [câu hỏi x10], "normal": [...], "hard": [...]}
#   câu hỏi: {"q": "...___...", "options": {"A":..,"B":..,"C":..,"D":..}, "answer": "A", "explain": "..."}
# Được kiểm tra rồi ghi ra src/data/quizzes/gXX.json để trang chủ đề import.
# ---------------------------------------------------------------------------
QUIZ_SRC = ROOT / "content" / "quizzes"
QUIZ_OUT = DATA / "quizzes"
LEVELS = ("easy", "normal", "hard")


def load_quiz(key: str, code: str) -> dict | None:
    path = QUIZ_SRC / f"{key}.json"
    if not path.exists():
        return None
    data = json.loads(path.read_text(encoding="utf-8"))
    out, errors, seen = {}, [], set()
    for level in LEVELS:
        items = data.get(level, [])
        if len(items) != 10:
            errors.append(f"{level}: cần 10 câu, có {len(items)}")
        out[level] = []
        for i, q in enumerate(items, start=1):
            where = f"{level} #{i}"
            opts = q.get("options", {})
            if sorted(opts) != list("ABCD") or not all(str(v).strip() for v in opts.values()):
                errors.append(f"{where}: options phải có đủ A-D")
            elif len({str(v).strip().lower() for v in opts.values()}) != 4:
                errors.append(f"{where}: có 2 phương án trùng nhau")
            if q.get("answer") not in "ABCD" or len(q.get("answer", "")) != 1:
                errors.append(f"{where}: answer phải là A/B/C/D")
            if not q.get("q", "").strip() or not q.get("explain", "").strip():
                errors.append(f"{where}: thiếu câu hỏi hoặc giải thích")
            key = q.get("q", "").strip().lower()
            if key in seen:
                errors.append(f"{where}: câu hỏi bị lặp")
            seen.add(key)
            out[level].append({"no": i, "code": code, "q": q.get("q", ""),
                               "options": {k: str(v) for k, v in opts.items()},
                               "answer": q.get("answer"), "explain": q.get("explain", "")})
    if errors:
        raise SystemExit(f"Lỗi trong {path.name}:\n  " + "\n  ".join(errors))
    return out


QUIZ_OUT.mkdir(parents=True, exist_ok=True)
for f in QUIZ_OUT.glob("*.json"):
    f.unlink()
level_quizzes = {}
for pg in pages:
    quiz = load_quiz(pg["key"], pg["code"])
    if quiz:
        level_quizzes[pg["key"]] = quiz
        (QUIZ_OUT / f"{pg['key']}.json").write_text(
            json.dumps(quiz, ensure_ascii=False, indent=1), encoding="utf-8")


group_pos = {g["key"]: i for i, g in enumerate(GROUPS)}
for gi, g in enumerate(GROUPS, start=1):
    category(GEN_DIR / g["dir"], g["label"], gi + 1, g["desc"], collapsed=gi != 2)
    for si, sub in enumerate(g.get("sub", []), start=1):
        category(GEN_DIR / g["dir"] / sub["dir"], sub["label"], si + 1, collapsed=False)

for t in pages:
    g = GROUPS[group_pos[t["group"]]]
    sub = next((s for s in g.get("sub", []) if t["code"] in s["codes"]), None)
    siblings = [x["key"] for x in pages if x["group"] == t["group"] and (not sub or x["code"] in sub["codes"])]
    order = siblings.index(t["key"]) + 1
    day = t["day"]
    day_info = next((d for d in roadmap if d["day"] == day), None)

    parts = [f"""---
title: "{t['title']}"
sidebar_label: "{t['label']} · {t['short']}"
sidebar_position: {order + (1 if g['key'] == 'Thì' and not sub else 0)}
slug: /ngu-phap/{t['slug']}
description: "{t['title']} – cấu trúc, cách dùng, dấu hiệu nhận biết, ví dụ và lỗi thường gặp."
---

import TopicMeta from '@site/src/components/TopicMeta';

# {t['title']}

<TopicMeta code="{t['code']}" label="{t['label']}" />
"""]
    extra = load_enrichment(t["key"])
    counter = iter(range(1, 20))

    def h2(title):
        return f"## {next(counter)}. {title}"

    def section(title, body, key):
        return "\n\n".join(x for x in (h2(title), body, extra.get(key)) if x)

    if "intro" in extra:
        parts.append(extra["intro"])
    parts.append(section("Cấu trúc", structure_block(t["structure"]), "structure"))
    if t["usage"] or "usage" in extra:
        parts.append(section("Cách dùng", labeled_list(t["usage"]), "usage"))
    parts.append(section("Dấu hiệu nhận biết", (
        labeled_list(t["signals"]) if t["signals"] else "" if "signals" in extra
        else "_Chủ đề này không có từ khóa nhận biết cố định – hãy dựa vào **cấu trúc** và **ngữ cảnh** của câu._"),
        "signals"))
    parts.append(section("Ví dụ", "\n".join(f"- {esc(e)}" for e in t["examples"]), "examples"))
    parts.append(section("Lỗi thường gặp", mistakes_block(t["mistakes"]), "mistakes"))
    if "compare" in extra:
        parts.append(f"{h2('Phân biệt & mở rộng')}\n\n{extra['compare']}")
    if "summary" in extra:
        parts.append(f"{h2('Tóm tắt ghi nhớ')}\n\n{extra['summary']}")
    if day_info:
        tasks = "\n".join(f"- {esc(x)}" for x in day_info["tasks"])
        parts.append(h2(f"Bài tập trong ngày (Ngày {day} · {day_info['minutes']} phút)") + "\n\n"
                     f"Chủ đề này nằm ở **[Ngày {day}](/docs/lo-trinh/{slugify(day_info['week'])}#ngay-{day})** "
                     f"của lộ trình. Chia 60 phút như sau:\n\n{tasks}")
    if t["key"] in level_quizzes:
        n = sum(len(v) for v in level_quizzes[t["key"]].values())
        var = f"quiz{t['key'].upper()}"
        parts.append(h2(f"Bài kiểm tra theo cấp độ ({n} câu)") + "\n\n"
                     "Ba cấp độ **Dễ – Trung bình – Khó**, mỗi cấp 10 câu. Nên làm lần lượt từ Dễ đến Khó; "
                     "đạt ≥ 80% ở một cấp rồi mới chuyển sang cấp tiếp theo. Điểm cao nhất được lưu trên tab.\n\n"
                     f"import {var} from '@site/src/data/quizzes/{t['key']}.json';\n"
                     "import LevelQuiz from '@site/src/components/LevelQuiz';\n\n"
                     f"<LevelQuiz topic=\"{t['label']}\" data={{{var}}} />")
    write(DOCS / t["dir"] / f"{t['slug']}.mdx", "\n\n".join(parts))

# ---------------------------------------------------------------------------
# Lộ trình theo tuần
# ---------------------------------------------------------------------------
LO = DOCS / "lo-trinh"
for f in LO.glob("tuan-*.mdx"):
    f.unlink()

week_titles = {
    "Tuần 1": "Nền tảng & các thì cơ bản",
    "Tuần 2": "Thì nâng cao, động từ khuyết thiếu, V-ing/to V",
    "Tuần 3": "Danh từ, tính từ, so sánh, giới từ, bị động",
    "Tuần 4": "Câu phức & cấu trúc nâng cao + Tổng kết",
}
test_by_day = {t["day"]: t for t in tests}


def topic_links(codes):
    return ", ".join(f"[{pg['label']} · {pg['short']}]({pg['path']})" for c in codes for pg in pages_by_code.get(c, []))


for wi, (week, title) in enumerate(week_titles.items(), start=1):
    days = [d for d in roadmap if d["week"] == week or (week == "Tuần 4" and d["week"] == "Tổng kết")]
    parts = [f"""---
title: "{week}: {title}"
sidebar_label: "{week} · Ngày {days[0]['day']}–{days[-1]['day']}"
sidebar_position: {wi + 1}
---

import DayCheck from '@site/src/components/DayCheck';

# {week}: {title}

| Ngày | Chủ đề | Tham chiếu |
|---|---|---|
""" + "\n".join(f"| [{d['day']}](#ngay-{d['day']}) | {esc_cell(d['topic'])} | {d['ref']} |" for d in days)]
    for d in days:
        body = [f"## Ngày {d['day']}: {esc(d['topic'])} {{#ngay-{d['day']}}}",
                f"<DayCheck day={{{d['day']}}} />"]
        if d["isTest"]:
            test = test_by_day.get(d["day"])
            body.append(f":::info[Ngày kiểm tra]\n\nÔn lại các chủ đề {d['ref']}, sau đó làm "
                        f"**[{test['name']}](/docs/bai-kiem-tra/{test['id']})** "
                        f"({len(test['questions'])} câu, không xem tài liệu). Mục tiêu **≥ 80%**.\n\n:::")
        else:
            body.append(f"**Chủ đề cần học:** {topic_links(d['codes'])}")
        body.append("**Kế hoạch 60 phút:**\n\n" + "\n".join(f"- {esc(x)}" for x in d["tasks"]))
        parts.append("\n\n".join(body))
    write(LO / f"tuan-{wi}.mdx", "\n\n".join(parts))

# ---------------------------------------------------------------------------
# Bài kiểm tra
# ---------------------------------------------------------------------------
KT = DOCS / "bai-kiem-tra"
for i, test in enumerate(tests, start=1):
    codes = sorted({q["code"] for q in test["questions"]})
    level_count = {lv: sum(q["level"] == lv for q in test["questions"]) for lv in ("basic", "normal", "hard")}
    parts_table = "\n".join(
        f"| {label} | {level_count[lv]} | {desc} |"
        for lv, label, desc in [
            ("basic", "🔵 Cơ bản", "Nhận biết công thức, dấu hiệu quen thuộc"),
            ("normal", "🟡 Trung bình", "Phủ định, câu hỏi, phân biệt các dạng dễ nhầm, dựa vào ngữ cảnh"),
            ("hard", "🔴 Khó", "Ngoại lệ, tìm lỗi sai, chọn câu đồng nghĩa – dạng đề thi"),
        ] if level_count[lv])
    write(KT / f"{test['id']}.mdx", f"""---
title: "{test['name']}: {test['subtitle']}"
sidebar_label: "{test['name']}"
sidebar_position: {i + 1}
slug: /bai-kiem-tra/{test['id']}
---

import Quiz from '@site/src/components/Quiz';

# {test['name']}

**{test['subtitle']}** · Làm vào **Ngày {test['day']}** · **{len(test['questions'])} câu** trắc nghiệm · Mục tiêu **≥ 80%**

| Phần | Số câu | Nội dung |
|---|---|---|
{parts_table}

:::note[Cách làm]
1. Làm hết các câu **không xem tài liệu**, theo thứ tự Cơ bản → Trung bình → Khó, rồi bấm **Nộp bài**.
2. Xem điểm tổng, điểm từng phần, đáp án đúng và giải thích cho từng câu.
3. Dưới 80%: mở các chủ đề trong mục **Cần ôn lại** rồi làm lại bài.
:::

<details>
<summary>Chủ đề trong bài ({len(codes)})</summary>

{topic_links(codes)}

</details>

<Quiz mode="exam" testId="{test['id']}" />
""")

# ---------------------------------------------------------------------------
# Tra cứu
# ---------------------------------------------------------------------------
TC = DOCS / "tra-cuu"
rows = []
for g in GROUPS:
    for t in (x for x in pages if x["group"] == g["key"]):
        first = t["structure"][0] if t["structure"] else ""
        rows.append(f"| {t['label']} | {g['label'][3:]} | [{esc_cell(t['short'])}]({t['path']}) | `{first.replace('|', '/')}` | {t['day'] or ''} |")
write(TC / "bang-tra-cuu.mdx", f"""---
title: Bảng tra cứu nhanh 29 chủ đề
sidebar_label: Bảng tra cứu nhanh
sidebar_position: 1
---

# Bảng tra cứu nhanh

Toàn bộ 29 chủ đề với công thức cốt lõi. Bấm vào tên chủ đề để xem chi tiết.

| Mã | Nhóm | Chủ đề | Công thức chính | Ngày học |
|---|---|---|---|---|
""" + "\n".join(rows))

mk = []
for g in GROUPS:
    members = [t for t in pages if t["group"] == g["key"] and t["mistakes"]]
    mk.append(f"## {g['label']}")
    for t in members:
        mk.append(f"### [{t['label']} · {esc(t['short'])}]({t['path']})\n\n" + mistakes_block(t["mistakes"]))
write(TC / "loi-thuong-gap.mdx", """---
title: Sổ tay lỗi thường gặp
sidebar_label: Sổ tay lỗi thường gặp
sidebar_position: 2
---

# Sổ tay lỗi thường gặp

Tổng hợp mọi lỗi sai hay gặp của 29 chủ đề. **Đọc lại trang này trước mỗi bài kiểm tra.**

""" + "\n\n".join(mk))

print(f"OK: {len(topics)} chủ đề, {len(roadmap)} ngày, {len(tests)} bài test, "
      f"{sum(len(t['questions']) for t in tests)} câu hỏi.")
